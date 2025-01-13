    
def read_pdf_pages(file_path):
    """Read the pdf file and return a list of strings on each page of the pdf"""
    """Read PDF file and return a list of strings for each page"""
    import fitz
    doc = fitz.open(file_path)
    documents = []
    for page in doc:
        documents.append(page.get_text())
    return documents
    
def read_word_pages(file_path):
    """Read the word file and return a list of word paragraph strings"""
    """Read Word file and return a list of paragraph strings"""
    # https://zhuanlan.zhihu.com/p/146363527
    from docx import Document
    # Open document
    document = Document(file_path)
    # Read titles, paragraphs, and list content
    ps = [ paragraph.text for paragraph in document.paragraphs]
    return ps

def read_ppt(file_path):
    import pptx
    prs = pptx.Presentation(file_path)
    documents = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                documents.append(shape.text)
    return '\n'.join(documents)


def read_file_content(file_path):
    """return content of txt, md, pdf, docx file"""
    # Supported file_path types include txt, md, pdf, docx
    if file_path.endswith('.pdf'):
        return ' '.join(read_pdf_pages(file_path))
    elif file_path.endswith('.docx'):
        return ' '.join(read_word_pages(file_path))
    elif file_path.endswith('.ppt') or file_path.endswith('.pptx'):
        return read_ppt(file_path)
    else:
        # Treat as text file by default
        with open(file_path, 'r', encoding='utf-8') as f:
            return '\n'.join(f.readlines())

def write_file_content(file_path, content):
    """write content to txt, md"""
    with open(file_path, 'w', encoding='utf-8') as f:
        f.write(content)
